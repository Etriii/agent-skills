#!/usr/bin/env python3
r"""
md_to_docx.py – turn a QA-template markdown document into a consistently styled Word file.

Usage
  python md_to_docx.py recommendations/02_QA_Audit_Report.md            # → 02_QA_Audit_Report.docx next to it
  python md_to_docx.py file.md -o /somewhere/else/name.docx              # choose the output path
  python md_to_docx.py --all recommendations/                            # convert every NN_*.md in the folder
  python md_to_docx.py --all v1/recommendations/md --outdir v1/recommendations/docs   # the md/ -> docs/ split
  python md_to_docx.py file.md --page letter                             # US Letter instead of A4

Needs:  pip install python-docx

What it understands (see 03_OUTPUT_TEMPLATES.md, part B):
  front matter (--- key: value ---) → cover block          headings # ## ### ####
  paragraphs with **bold** *italic* `code` [text](url)     - bullets / 1. numbered (one nesting level, 2-space indent)
  | pipe | tables |  (empty header row = no header)         > blockquote → shaded note
  ``` fenced code ```                                       --- → horizontal rule      \pagebreak → page break
  <br> inside table cells → line break
  ![caption](path/to/image.png) on its own line → picture, scaled to the page width, caption in small grey text
                                                   (path relative to the .md file, e.g. ../evidence/F-003_invoice.png)
  <!-- widths: 1.2, 3, 7.5, 2.3, 3 -->  on the line before a table → column widths in cm (scaled to fit the page)
  Status / severity / priority words in table cells are colour-coded automatically.

  python md_to_docx.py --version   prints the kit version (from the VERSION file next to tools/)

Everything about the look (fonts, colours, margins, header/footer text) is in the SETTINGS block below.
"""

import argparse
import datetime as _dt
import os
import re
import sys

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Emu, Pt, RGBColor
except ImportError:  # pragma: no cover
    sys.exit("python-docx is not installed. Run:  pip install python-docx")

# ─────────────────────────────────────────────────────────────────────────────
# SETTINGS – change the look here
# ─────────────────────────────────────────────────────────────────────────────
FONT_BODY = "Calibri"
FONT_CODE = "Consolas"
SIZE_BODY = 11
SIZE_TABLE = 9.5
SIZE_CODE = 9
COLOR_HEADING = RGBColor(0x1F, 0x38, 0x64)   # dark navy
COLOR_MUTED = RGBColor(0x59, 0x59, 0x59)     # grey for header/footer/meta
COLOR_RULE = "A6A6A6"                        # horizontal rule
FILL_TABLE_HEADER = "DEE6F0"                 # light blue-grey
FILL_NOTE = "F2F2F2"                         # blockquote box
FILL_CODE = "F5F5F5"
MARGIN_CM = 2.0

# Words that get a coloured cell when they stand alone (or start) a table cell.
# (key, match mode, fill, text colour)   mode: "prefix" = cell starts with it, "exact" = whole cell
STATUS_COLOURS = [
    ("PASS",         "prefix", "D9EAD3", "274E13"),
    ("FAIL",         "prefix", "F4CCCC", "990000"),
    ("MISSING",      "prefix", "FCE5CD", "B45F06"),
    ("N/A",          "prefix", "EFEFEF", "595959"),
    ("NEEDS-MANUAL", "prefix", "FFF2CC", "7F6000"),
    ("NOT CHECKED",  "prefix", "FFF2CC", "7F6000"),
    ("FIXED",        "prefix", "CFE2F3", "0B5394"),
    ("CRITICAL",     "exact",  "F4CCCC", "990000"),
    ("HIGH",         "exact",  "FCE5CD", "B45F06"),
    ("MEDIUM",       "exact",  "FFF2CC", "7F6000"),
    ("LOW",          "exact",  "EFEFEF", "595959"),
    ("MUST",         "exact",  "F4CCCC", "990000"),
    ("SHOULD",       "exact",  "FCE5CD", "B45F06"),
    ("NICE",         "exact",  "EFEFEF", "595959"),
]

def kit_version():
    """Version of the kit this script belongs to, read from the VERSION file beside the tools folder."""
    here = os.path.dirname(os.path.abspath(__file__))
    for candidate in (os.path.join(here, "..", "VERSION"), os.path.join(here, "VERSION")):
        try:
            with open(candidate, encoding="utf-8") as f:
                return f.read().strip() or "unknown"
        except OSError:
            continue
    return "unknown"


# Cover block: which front-matter keys to show, in this order, with their labels.
COVER_FIELDS = [
    ("project", "Project"),
    ("document", "Document"),
    ("date", "Date"),
    ("reviewer", "Reviewer"),
    ("version", "Version"),
    ("supersedes", "Supersedes"),
    ("status", "Status"),
    ("reviewed", "What was reviewed"),
]

# ─────────────────────────────────────────────────────────────────────────────
# Low-level XML helpers (keep child elements in schema order so Word never complains)
# ─────────────────────────────────────────────────────────────────────────────
_PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
              "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
              "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
              "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
              "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
              "sectPr", "pPrChange"]
_TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
               "textDirection", "tcFitText", "vAlign", "hideMark"]
_TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
                "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
                "tblCaption", "tblDescription"]


def _insert_in_order(parent, elem, order):
    tag = elem.tag.split("}")[1]
    my_pos = order.index(tag) if tag in order else len(order)
    for child in list(parent):
        ctag = child.tag.split("}")[1]
        cpos = order.index(ctag) if ctag in order else len(order)
        if cpos > my_pos:
            child.addprevious(elem)
            return elem
    parent.append(elem)
    return elem


def _shd(fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    return shd


def shade_paragraph(p, fill):
    _insert_in_order(p._p.get_or_add_pPr(), _shd(fill), _PPR_ORDER)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    _insert_in_order(tcPr, _shd(fill), _TCPR_ORDER)


def paragraph_border(p, sides=("bottom",), color=COLOR_RULE, sz=6, space=1):
    pBdr = OxmlElement("w:pBdr")
    for side in sides:
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(sz))
        b.set(qn("w:space"), str(space))
        b.set(qn("w:color"), color)
        pBdr.append(b)
    _insert_in_order(p._p.get_or_add_pPr(), pBdr, _PPR_ORDER)


def add_field(run, instr):
    """Insert a Word field (PAGE, NUMPAGES …) into a run."""
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")  # ask the viewer to recompute the field rather than trust the cached value
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr_el, sep, txt, end):
        run._r.append(el)


def set_cell_margins(table, top=40, bottom=40, left=80, right=80):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    _insert_in_order(tblPr, mar, _TBLPR_ORDER)


def set_table_widths(table, widths_cm):
    """Fixed layout with explicit widths so Word and LibreOffice render the same columns."""
    table.autofit = False  # python-docx writes <w:tblLayout w:type="fixed"/> for us
    grid = table._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        if i < len(widths_cm):
            gc.set(qn("w:w"), str(int(widths_cm[i] * 567)))  # cm → twips
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths_cm):
                cell.width = Cm(widths_cm[i])


def column_widths(header, rows, usable_cm, min_cm=1.5):
    """Share the usable width between columns according to how much text they hold."""
    ncols = max([len(header) if header else 0] + [len(r) for r in rows])
    texts = [[] for _ in range(ncols)]
    if header:
        for i, h in enumerate(header):
            texts[i].append(h)
    for r in rows:
        for i, c in enumerate(r):
            texts[i].append(c)
    scores, minimums = [], []
    for col in texts:
        plain = [re.sub(r"[*`_]", "", t) for t in col] or [""]
        lens = [len(t) for t in plain]
        longest = min(max(lens), 90)
        average = sum(lens) / len(lens)
        scores.append(max(4.0, 0.55 * longest + 0.45 * average + 2))
        # a column must at least fit its longest single word (capped, so file paths do not hog the table)
        longest_word = max((len(w) for t in plain for w in t.split()), default=0)
        cap = min(3.2, usable_cm / ncols * 1.4)  # with many columns, accept some mid-word wrapping
        minimums.append(max(min(min_cm, cap), min(cap, 0.2 * longest_word + 0.5)))
    total = sum(scores)
    widths = [usable_cm * s / total for s in scores]
    # enforce the minimums, taking the difference from the columns that have room
    for _ in range(ncols + 1):
        short = [i for i, w in enumerate(widths) if w < minimums[i]]
        if not short:
            break
        deficit = sum(minimums[i] - widths[i] for i in short)
        for i in short:
            widths[i] = minimums[i]
        room = [i for i, w in enumerate(widths) if w > minimums[i]]
        room_total = sum(widths[i] - minimums[i] for i in room) or 1
        for i in room:
            widths[i] -= deficit * (widths[i] - minimums[i]) / room_total
    # never overflow the page
    total = sum(widths)
    if total > usable_cm:
        widths = [w * usable_cm / total for w in widths]
    return widths


def repeat_header_row(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


# ─────────────────────────────────────────────────────────────────────────────
# Numbering (bullets share one list; every numbered list restarts at 1)
# ─────────────────────────────────────────────────────────────────────────────
class Numbering:
    BULLET_ABS = 900
    DECIMAL_ABS = 901

    def __init__(self, doc):
        self.doc = doc
        self.root = doc.part.numbering_part.element
        self._add_abstract(self.BULLET_ABS, [("bullet", "•"), ("bullet", "–")])
        self._add_abstract(self.DECIMAL_ABS, [("decimal", "%1."), ("lowerLetter", "%2.")])
        self.bullet_num = self._add_num(self.BULLET_ABS)

    def _next_num_id(self):
        ids = [int(n.get(qn("w:numId"))) for n in self.root.findall(qn("w:num"))]
        return (max(ids) + 1) if ids else 1

    def _add_abstract(self, abs_id, levels):
        absn = OxmlElement("w:abstractNum")
        absn.set(qn("w:abstractNumId"), str(abs_id))
        mlt = OxmlElement("w:multiLevelType")
        mlt.set(qn("w:val"), "hybridMultilevel")
        absn.append(mlt)
        for i, (fmt, text) in enumerate(levels):
            lvl = OxmlElement("w:lvl")
            lvl.set(qn("w:ilvl"), str(i))
            start = OxmlElement("w:start"); start.set(qn("w:val"), "1"); lvl.append(start)
            nf = OxmlElement("w:numFmt"); nf.set(qn("w:val"), fmt); lvl.append(nf)
            lt = OxmlElement("w:lvlText"); lt.set(qn("w:val"), text); lvl.append(lt)
            jc = OxmlElement("w:lvlJc"); jc.set(qn("w:val"), "left"); lvl.append(jc)
            pPr = OxmlElement("w:pPr")
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str(720 + 360 * i))
            ind.set(qn("w:hanging"), "360")
            pPr.append(ind)
            lvl.append(pPr)
            absn.append(lvl)
        first_num = self.root.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(absn)
        else:
            self.root.append(absn)

    def _add_num(self, abs_id):
        num_id = self._next_num_id()
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        a = OxmlElement("w:abstractNumId"); a.set(qn("w:val"), str(abs_id)); num.append(a)
        ov = OxmlElement("w:lvlOverride"); ov.set(qn("w:ilvl"), "0")
        so = OxmlElement("w:startOverride"); so.set(qn("w:val"), "1"); ov.append(so); num.append(ov)
        self.root.append(num)
        return num_id

    def new_numbered_list(self):
        return self._add_num(self.DECIMAL_ABS)

    @staticmethod
    def apply(paragraph, num_id, level):
        pPr = paragraph._p.get_or_add_pPr()
        numPr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), str(level)); numPr.append(ilvl)
        nid = OxmlElement("w:numId"); nid.set(qn("w:val"), str(num_id)); numPr.append(nid)
        _insert_in_order(pPr, numPr, _PPR_ORDER)


# ─────────────────────────────────────────────────────────────────────────────
# Inline markdown → runs
# ─────────────────────────────────────────────────────────────────────────────
_INLINE_RE = re.compile(
    r"(\*\*(?P<b>.+?)\*\*)"          # **bold**
    r"|(`(?P<c>[^`]+?)`)"            # `code`
    r"|(\[(?P<lt>[^\]]+?)\]\((?P<lu>[^)]+?)\))"   # [text](url)
    r"|((?<![\w*])\*(?P<i>[^*\n]+?)\*(?![\w*]))"   # *italic*
    r"|((?<!\w)_(?P<i2>[^_\n]+?)_(?!\w))"          # _italic_
    r"|(?P<br><br\s*/?>)"
)


_ESC_MAP = {"*": "\ue001", "`": "\ue002", "_": "\ue003", "|": "\ue004", "[": "\ue005", "]": "\ue006",
            "<": "\ue007", ">": "\ue008", "\\": "\ue009"}
_ESC_BACK = {v: k for k, v in _ESC_MAP.items()}


def _protect_escapes(text):
    r"""Turn backslash-escaped characters (\* \` \_ \| \[ \] \< \> \\) into private-use characters so the inline parser ignores them."""
    def repl(m):
        ch = m.group(1)
        return _ESC_MAP.get(ch, ch)
    return re.sub(r"\\([\\`*_{}\[\]()#+\-.!|>~<])", repl, text)


def _unescape(text):
    for k, v in _ESC_BACK.items():
        text = text.replace(k, v)
    return text


def add_inline(paragraph, text, base_bold=False, base_italic=False, size=None, color=None, _protected=False):
    """Write markdown-formatted text into a paragraph as runs."""
    if not _protected:
        text = _protect_escapes(text)
    pos = 0
    for m in _INLINE_RE.finditer(text):
        if m.start() > pos:
            _run(paragraph, text[pos:m.start()], base_bold, base_italic, size=size, color=color)
        if m.group("b") is not None:
            add_inline(paragraph, m.group("b"), True, base_italic, size, color, _protected=True)
        elif m.group("c") is not None:
            _run(paragraph, m.group("c"), base_bold, base_italic, code=True, size=size)
        elif m.group("lt") is not None:
            _run(paragraph, m.group("lt"), base_bold, base_italic, size=size, color=RGBColor(0x0B, 0x53, 0x94), underline=True)
            _run(paragraph, f" ({m.group('lu')})", base_bold, base_italic, size=size or SIZE_BODY, color=COLOR_MUTED)
        elif m.group("i") is not None:
            add_inline(paragraph, m.group("i"), base_bold, True, size, color, _protected=True)
        elif m.group("i2") is not None:
            add_inline(paragraph, m.group("i2"), base_bold, True, size, color, _protected=True)
        elif m.group("br") is not None:
            paragraph.add_run().add_break(WD_BREAK.LINE)
        pos = m.end()
    if pos < len(text):
        _run(paragraph, text[pos:], base_bold, base_italic, size=size, color=color)


def _run(paragraph, text, bold=False, italic=False, code=False, size=None, color=None, underline=False):
    text = _unescape(text)
    if not text:
        return None
    r = paragraph.add_run(text)
    r.bold = bold or None
    r.italic = italic or None
    if underline:
        r.underline = True
    if code:
        r.font.name = FONT_CODE
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CODE)
        r.font.size = Pt((size or SIZE_BODY) - 1.5)
        r.font.color.rgb = RGBColor(0x7A, 0x1F, 0x4F)
    else:
        if size:
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
    return r


# ─────────────────────────────────────────────────────────────────────────────
# Block parsing
# ─────────────────────────────────────────────────────────────────────────────
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_BULLET_RE = re.compile(r"^(\s*)[-*+]\s+(.*)$")
_NUMBER_RE = re.compile(r"^(\s*)(\d+)[.)]\s+(.*)$")
_TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_RULE_RE = re.compile(r"^\s*([-*_])(\s*\1){2,}\s*$")
_PAGEBREAK_RE = re.compile(r"^\s*(\\pagebreak|<!--\s*pagebreak\s*-->)\s*$", re.I)
_IMAGE_RE = re.compile(r"^\s*!\[(?P<alt>[^\]]*)\]\((?P<src>[^)\s]+)(?:\s+\"[^\"]*\")?\)\s*$")
_WIDTHS_RE = re.compile(r"^\s*<!--\s*widths?\s*:\s*([\d.\s,]+?)\s*-->\s*$", re.I)


def parse_front_matter(lines):
    meta = {}
    if lines and lines[0].strip() == "---":
        for i in range(1, len(lines)):
            if lines[i].strip() == "---":
                for ln in lines[1:i]:
                    if ":" in ln:
                        k, v = ln.split(":", 1)
                        meta[k.strip().lower()] = v.strip()
                return meta, lines[i + 1:]
    return meta, lines


def split_cells(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|") and not line.endswith("\\|"):
        line = line[:-1]
    parts = re.split(r"(?<!\\)\|", line)
    return [p.strip().replace("\\|", "|") for p in parts]


class Converter:
    def __init__(self, page="a4"):
        self.doc = Document()
        self.page = page
        self._setup_styles()
        self._setup_page()
        zoom = self.doc.settings.element.find(qn("w:zoom"))
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")
        self.num = Numbering(self.doc)
        self.meta = {}
        self.first_h1_skipped = False
        self.base_dir = None          # folder of the source .md, for resolving image paths
        self.pending_widths = None    # column widths (cm) declared by a <!-- widths: --> hint for the next table

    # ---- document set-up -------------------------------------------------------------------
    def _setup_styles(self):
        st = self.doc.styles
        normal = st["Normal"]
        normal.font.name = FONT_BODY
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
        normal.font.size = Pt(SIZE_BODY)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.08
        for name, size, bold, italic, before, after in (
            ("Title", 26, True, False, 0, 4),
            ("Subtitle", 14, False, False, 0, 10),
            ("Heading 1", 17, True, False, 18, 6),
            ("Heading 2", 13.5, True, False, 14, 4),
            ("Heading 3", 12, True, False, 10, 3),
            ("Heading 4", 11, True, True, 8, 2),
        ):
            s = st[name]
            s.font.name = FONT_BODY
            s.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
            s.font.size = Pt(size)
            s.font.bold = bold
            s.font.italic = italic
            s.font.color.rgb = COLOR_HEADING if name != "Subtitle" else COLOR_MUTED
            s.paragraph_format.space_before = Pt(before)
            s.paragraph_format.space_after = Pt(after)
            s.paragraph_format.keep_with_next = True
        # Word's built-in Title style has a bottom border and odd spacing – neutralise it.
        title_pPr = st["Title"].element.get_or_add_pPr()
        for bdr in title_pPr.findall(qn("w:pBdr")):
            title_pPr.remove(bdr)

    def _setup_page(self):
        sec = self.doc.sections[0]
        if self.page == "letter":
            sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
        else:
            sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.orientation = WD_ORIENT.PORTRAIT
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Cm(MARGIN_CM))
        sec.header_distance = Cm(1.0)
        sec.footer_distance = Cm(1.0)
        self.usable_cm = (sec.page_width - sec.left_margin - sec.right_margin) / 360000  # EMU → cm

    def _header_footer(self, source_name):
        sec = self.doc.sections[0]
        project = self.meta.get("project", "")
        document = self.meta.get("document", "")
        hp = sec.header.paragraphs[0]
        hp.text = ""
        r = hp.add_run(project)
        r.font.size = Pt(9); r.font.color.rgb = COLOR_MUTED
        hp.add_run("\t")
        r = hp.add_run(document)
        r.font.size = Pt(9); r.font.color.rgb = COLOR_MUTED; r.bold = True
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        # right-aligned tab stop for the document name
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right")
        usable = sec.page_width - sec.left_margin - sec.right_margin
        tab.set(qn("w:pos"), str(int(usable / 635)))  # EMU → twips
        tabs.append(tab)
        _insert_in_order(hp._p.get_or_add_pPr(), tabs, _PPR_ORDER)
        paragraph_border(hp, ("bottom",), color=COLOR_RULE, sz=4)

        fp = sec.footer.paragraphs[0]
        fp.text = ""
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("Page "); r.font.size = Pt(9); r.font.color.rgb = COLOR_MUTED
        r = fp.add_run(); r.font.size = Pt(9); r.font.color.rgb = COLOR_MUTED; add_field(r, "PAGE")
        r = fp.add_run(" of "); r.font.size = Pt(9); r.font.color.rgb = COLOR_MUTED
        r = fp.add_run(); r.font.size = Pt(9); r.font.color.rgb = COLOR_MUTED; add_field(r, "NUMPAGES")
        stamp = f"   ·   {self.meta.get('status', '')}   ·   built from {source_name} on {_dt.date.today().isoformat()}"
        r = fp.add_run(stamp); r.font.size = Pt(8); r.font.color.rgb = COLOR_MUTED

    def _core_properties(self):
        cp = self.doc.core_properties
        cp.title = self.meta.get("document", "")
        cp.subject = self.meta.get("project", "")
        cp.author = self.meta.get("reviewer", "QA review")
        cp.comments = f"Generated by the Ultimate QA Template {kit_version()} (md_to_docx.py)"

    # ---- cover block -----------------------------------------------------------------------
    def _cover(self):
        m = self.meta
        if not m:
            return
        p = self.doc.add_paragraph(style="Title")
        p.add_run(m.get("document", "Document"))
        if m.get("project"):
            p = self.doc.add_paragraph(style="Subtitle")
            p.add_run(m["project"])
        rows = [(label, m[key]) for key, label in COVER_FIELDS if m.get(key)]
        for key in m:
            if key not in dict(COVER_FIELDS) and key not in ("title",):
                rows.append((key.replace("_", " ").capitalize(), m[key]))
        if rows:
            t = self.doc.add_table(rows=0, cols=2)
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.LEFT
            set_cell_margins(t)
            for label, value in rows:
                cells = t.add_row().cells
                pr = cells[0].paragraphs[0]
                _run(pr, label, bold=True, size=SIZE_TABLE, color=COLOR_HEADING)
                shade_cell(cells[0], FILL_TABLE_HEADER)
                pr = cells[1].paragraphs[0]
                add_inline(pr, value, size=SIZE_TABLE)
                if label == "Status":
                    self._colour_status_cell(cells[1], value)
            set_table_widths(t, [4.0, self.usable_cm - 4.0])
            self._tighten_table(t)
        rule = self.doc.add_paragraph()
        paragraph_border(rule, ("bottom",), color=COLOR_RULE, sz=8)
        rule.paragraph_format.space_after = Pt(10)

    # ---- tables ----------------------------------------------------------------------------
    def _colour_status_cell(self, cell, text):
        t = text.strip()
        up = t.upper()
        for key, mode, fill, colour in STATUS_COLOURS:
            hit = (up == key) if mode == "exact" else (up == key or up.startswith(key + " ") or up.startswith(key + "(") or up.startswith(key + "–") or up.startswith(key + "-") and key != "N/A")
            if mode == "exact" and not hit:
                # allow "High (…)" style
                hit = up.startswith(key + " (")
            if hit:
                shade_cell(cell, fill)
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.color.rgb = RGBColor.from_string(colour)
                        r.bold = True
                return

    @staticmethod
    def _tighten_table(t):
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.line_spacing = 1.0

    def add_table(self, header, rows):
        ncols = max([len(header) if header else 0] + [len(r) for r in rows])
        if ncols == 0:
            return
        t = self.doc.add_table(rows=0, cols=ncols)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        set_cell_margins(t)
        widths = column_widths(header, rows, self.usable_cm)
        # key/value block (no header, two columns): fixed narrow label column
        if not (header and any(h.strip() for h in header)) and ncols == 2:
            widths = [3.6, self.usable_cm - 3.6]
        # an explicit <!-- widths: --> hint wins, scaled so the table still fits the page
        if self.pending_widths:
            hint = [w for w in self.pending_widths if w > 0]
            if len(hint) == ncols:
                scale = self.usable_cm / sum(hint)
                widths = [w * scale for w in hint]
            self.pending_widths = None
        if header and any(h.strip() for h in header):
            cells = t.add_row().cells
            for i in range(ncols):
                txt = header[i] if i < len(header) else ""
                add_inline(cells[i].paragraphs[0], txt, base_bold=True, size=SIZE_TABLE, color=COLOR_HEADING)
                shade_cell(cells[i], FILL_TABLE_HEADER)
            repeat_header_row(t.rows[0])
        for r in rows:
            cells = t.add_row().cells
            for i in range(ncols):
                txt = r[i] if i < len(r) else ""
                add_inline(cells[i].paragraphs[0], txt, size=SIZE_TABLE)
                plain = re.sub(r"[*`_]", "", txt).strip()
                self._colour_status_cell(cells[i], plain)
        set_table_widths(t, widths)
        self._tighten_table(t)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(2)

    # ---- other blocks ----------------------------------------------------------------------
    def add_heading(self, level, text):
        if level == 1 and not self.first_h1_skipped:
            self.first_h1_skipped = True
            if self.meta.get("document") and self._norm(text) == self._norm(self.meta["document"]):
                return  # already shown on the cover
        style = f"Heading {min(level, 4)}"
        p = self.doc.add_paragraph(style=style)
        add_inline(p, text)

    @staticmethod
    def _norm(s):
        return re.sub(r"[\s–—-]+", " ", s).strip().lower()

    def add_paragraph(self, text):
        p = self.doc.add_paragraph()
        add_inline(p, text)

    def add_list_item(self, text, level, numbered, num_id):
        p = self.doc.add_paragraph(style="List Paragraph")
        Numbering.apply(p, num_id if numbered else self.num.bullet_num, level)
        p.paragraph_format.space_after = Pt(2)
        add_inline(p, text)

    def add_quote(self, lines):
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.right_indent = Cm(0.5)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        shade_paragraph(p, FILL_NOTE)
        paragraph_border(p, ("left",), color="1F3864", sz=18, space=4)
        for i, ln in enumerate(lines):
            if i:
                p.add_run().add_break(WD_BREAK.LINE)
            add_inline(p, ln)

    def add_code(self, lines):
        for i, ln in enumerate(lines):
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.left_indent = Cm(0.3)
            shade_paragraph(p, FILL_CODE)
            r = p.add_run(ln if ln else " ")
            r.font.name = FONT_CODE
            r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CODE)
            r.font.size = Pt(SIZE_CODE)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def add_image(self, src, alt):
        """Embed a picture from a path relative to the source .md, scaled down to the page width, with an optional caption."""
        path = src if os.path.isabs(src) else os.path.join(self.base_dir or os.getcwd(), src)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        if not os.path.isfile(path):
            r = p.add_run(f"[image not found: {src}]")
            r.italic = True
            r.font.size = Pt(SIZE_TABLE)
            r.font.color.rgb = COLOR_MUTED
            return
        try:
            pic = p.add_run().add_picture(path)
        except Exception as exc:  # unsupported format, corrupt file
            r = p.add_run(f"[image could not be embedded: {src} – {exc}]")
            r.italic = True
            r.font.size = Pt(SIZE_TABLE)
            r.font.color.rgb = COLOR_MUTED
            return
        max_w = Cm(self.usable_cm)
        if pic.width > max_w:
            ratio = max_w / pic.width
            pic.width = Emu(int(max_w))
            pic.height = Emu(int(pic.height * ratio))
        if alt.strip():
            cap = self.doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_after = Pt(8)
            add_inline(cap, alt.strip(), base_italic=True, size=SIZE_TABLE, color=COLOR_MUTED)

    def add_rule(self):
        p = self.doc.add_paragraph()
        paragraph_border(p, ("bottom",), color=COLOR_RULE, sz=6)
        p.paragraph_format.space_after = Pt(6)

    def add_pagebreak(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- main loop -------------------------------------------------------------------------
    def convert(self, md_text, source_name="document.md", base_dir=None):
        self.base_dir = base_dir or os.getcwd()
        lines = md_text.replace("\r\n", "\n").split("\n")
        self.meta, lines = parse_front_matter(lines)
        self._core_properties()
        self._header_footer(source_name)
        self._cover()

        i, n = 0, len(lines)
        para_buf = []
        list_state = {"numbered": None, "num_id": None}

        def flush_para():
            if para_buf:
                self.add_paragraph(" ".join(s.strip() for s in para_buf))
                para_buf.clear()

        def end_list():
            list_state["numbered"] = None
            list_state["num_id"] = None

        while i < n:
            ln = lines[i]
            stripped = ln.strip()

            # fenced code
            if stripped.startswith("```"):
                flush_para(); end_list()
                block = []
                i += 1
                while i < n and not lines[i].strip().startswith("```"):
                    block.append(lines[i].rstrip("\n"))
                    i += 1
                self.add_code(block)
                i += 1
                continue

            # page break / rule
            if _PAGEBREAK_RE.match(stripped):
                flush_para(); end_list(); self.add_pagebreak(); i += 1; continue
            if _RULE_RE.match(stripped):
                flush_para(); end_list(); self.add_rule(); i += 1; continue

            # column-width hint for the next table
            mw = _WIDTHS_RE.match(stripped)
            if mw:
                flush_para()
                try:
                    self.pending_widths = [float(x) for x in re.split(r"[,\s]+", mw.group(1).strip()) if x]
                except ValueError:
                    self.pending_widths = None
                i += 1
                continue

            # picture on its own line
            mi = _IMAGE_RE.match(stripped)
            if mi:
                flush_para(); end_list()
                self.add_image(mi.group("src"), mi.group("alt"))
                i += 1
                continue

            # heading
            m = _HEADING_RE.match(stripped)
            if m:
                flush_para(); end_list()
                self.add_heading(len(m.group(1)), m.group(2))
                i += 1
                continue

            # table
            if stripped.startswith("|"):
                flush_para(); end_list()
                block = []
                while i < n and lines[i].strip().startswith("|"):
                    block.append(lines[i].strip())
                    i += 1
                header, rows = None, []
                if len(block) >= 2 and _TABLE_SEP_RE.match(block[1]):
                    header = split_cells(block[0])
                    rows = [split_cells(b) for b in block[2:]]
                else:
                    rows = [split_cells(b) for b in block]
                self.add_table(header, rows)
                continue

            # blockquote
            if stripped.startswith(">"):
                flush_para(); end_list()
                block = []
                while i < n and lines[i].strip().startswith(">"):
                    block.append(lines[i].strip()[1:].strip())
                    i += 1
                self.add_quote(block)
                continue

            # lists
            mb, mn = _BULLET_RE.match(ln), _NUMBER_RE.match(ln)
            if (mb or mn) and not _RULE_RE.match(stripped):
                flush_para()
                indent = len((mb or mn).group(1).replace("\t", "    "))
                level = 1 if indent >= 2 else 0
                text = mb.group(2) if mb else mn.group(3)
                numbered = bool(mn)
                starts_over = numbered and level == 0 and mn.group(2) == "1" and list_state["numbered"] is True
                if level == 0 and (list_state["numbered"] is None or list_state["numbered"] != numbered or starts_over):
                    list_state["numbered"] = numbered
                    list_state["num_id"] = self.num.new_numbered_list() if numbered else self.num.bullet_num
                elif list_state["numbered"] is None:
                    list_state["numbered"] = numbered
                    list_state["num_id"] = self.num.new_numbered_list() if numbered else self.num.bullet_num
                # continuation lines belonging to this item (indented, not a new item)
                j = i + 1
                while j < n and lines[j].strip() and not (_BULLET_RE.match(lines[j]) or _NUMBER_RE.match(lines[j])) \
                        and (lines[j].startswith("  ") or lines[j].startswith("\t")) and not lines[j].strip().startswith(("|", "#", "```", ">")):
                    text += " " + lines[j].strip()
                    j += 1
                self.add_list_item(text, level, numbered if level == 0 else list_state["numbered"], list_state["num_id"])
                i = j
                continue

            # blank line
            if not stripped:
                flush_para()
                # a blank line only ends a list if the next non-blank line is not a list item
                k = i + 1
                while k < n and not lines[k].strip():
                    k += 1
                if k >= n or not (_BULLET_RE.match(lines[k]) or _NUMBER_RE.match(lines[k])):
                    end_list()
                i += 1
                continue

            # plain paragraph text
            end_list()
            para_buf.append(ln)
            i += 1
        flush_para()
        return self.doc


# ─────────────────────────────────────────────────────────────────────────────
# CLI
# ─────────────────────────────────────────────────────────────────────────────
def convert_file(src, dst=None, page="a4"):
    with open(src, encoding="utf-8") as f:
        md = f.read()
    conv = Converter(page=page)
    doc = conv.convert(md, source_name=os.path.basename(src), base_dir=os.path.dirname(os.path.abspath(src)))
    if dst is None:
        dst = os.path.splitext(src)[0] + ".docx"
    os.makedirs(os.path.dirname(os.path.abspath(dst)), exist_ok=True)
    doc.save(dst)
    return dst


def main(argv=None):
    ap = argparse.ArgumentParser(description="Convert QA-template markdown to a styled Word document.")
    ap.add_argument("inputs", nargs="*", help="markdown file(s)")
    ap.add_argument("-o", "--output", help="output .docx (only with a single input)")
    ap.add_argument("--all", metavar="FOLDER", help="convert every NN_*.md in FOLDER")
    ap.add_argument("--outdir", metavar="FOLDER",
                    help="write the .docx files into FOLDER instead of next to the .md "
                         "(used for the md/ -> docs/ split)")
    ap.add_argument("--page", choices=["a4", "letter"], default="a4")
    ap.add_argument("--version", action="version", version=f"md_to_docx.py (Ultimate QA Template {kit_version()})")
    args = ap.parse_args(argv)

    files = list(args.inputs)
    if args.all:
        for name in sorted(os.listdir(args.all)):
            if re.match(r"^\d{2}_.*\.md$", name):
                files.append(os.path.join(args.all, name))
    if not files:
        ap.error("give at least one .md file or --all FOLDER")
    if args.output and len(files) > 1:
        ap.error("-o can only be used with a single input file")
    if args.output and args.outdir:
        ap.error("use either -o or --outdir, not both")

    for src in files:
        dst = args.output
        if dst is None and args.outdir:
            base = os.path.splitext(os.path.basename(src))[0] + ".docx"
            dst = os.path.join(args.outdir, base)
        dst = convert_file(src, dst, page=args.page)
        print(f"wrote {dst}")


if __name__ == "__main__":
    main()
